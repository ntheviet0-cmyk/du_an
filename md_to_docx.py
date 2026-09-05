"""md_to_docx.py - Convert reports/ibm/REPORT.md to a formatted Word document."""
import re, os
from docx import Document
from docx.shared import Inches

SRC = "reports/ibm/REPORT.md"
OUT = "reports/ibm/REPORT_v2.docx"
IMG_DIR = "reports/ibm"

def strip_bold(s):
    return s.replace("**", "")

def add_rich(paragraph, text):
    # inline bold via **...**
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part == "":
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:
            run.bold = True

doc = Document()
lines = open(SRC, encoding="utf-8").read().split("\n")

def is_sep(row):
    return all(re.fullmatch(r"[-: ]+", c) for c in row)

i = 0
while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    # table block
    if stripped.startswith("|") and stripped.count("|") >= 2:
        tbl = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            tbl.append(lines[i]); i += 1
        rows = []
        for tl in tbl:
            rows.append([c.strip() for c in tl.strip().strip("|").split("|")])
        data = [r for r in rows if not is_sep(r)]
        if not data:
            continue
        ncol = len(data[0])
        t = doc.add_table(rows=1, cols=ncol)
        try:
            t.style = "Light Grid Accent 1"
        except Exception:
            pass
        for j, h in enumerate(data[0]):
            t.rows[0].cells[j].text = strip_bold(h)
        for r in data[1:]:
            cells = t.add_row().cells
            for j in range(min(len(r), ncol)):
                cells[j].text = strip_bold(r[j])
        doc.add_paragraph()
        continue
    # headings
    if stripped.startswith("#### "):
        doc.add_heading(stripped[5:], level=3)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=0)
    elif stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_rich(p, stripped[2:])
    elif stripped == "":
        pass
    else:
        txt = stripped
        if txt.startswith("*") and not txt.startswith("**"):
            txt = txt.strip("*")
        p = doc.add_paragraph()
        add_rich(p, txt)
    i += 1

# ---- Appendix: embed figures ----
doc.add_page_break()
doc.add_heading("Phu luc - Hinh anh minh hoa", level=1)
imgs = ["roc_curves.png", "metrics_bar.png", "attrition_dag.png",
        "shap_rf.png", "shap_xgboost.png"]
captions = {
    "roc_curves.png": "Hinh 1. Duong ROC: Random Forest vs XGBoost",
    "metrics_bar.png": "Hinh 2. Bien do so sanh cac chi so hieu nang",
    "attrition_dag.png": "Hinh 3. Do thi nhan qua (Causal DAG): OverTime -> Attrition",
    "shap_rf.png": "Hinh 4. SHAP summary - Random Forest",
    "shap_xgboost.png": "Hinh 5. SHAP summary - XGBoost",
}
for im in imgs:
    path = os.path.join(IMG_DIR, im)
    if os.path.exists(path):
        doc.add_paragraph(captions.get(im, im))
        doc.add_picture(path, width=Inches(6.0))
        doc.add_paragraph()

doc.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
